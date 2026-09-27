from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_ROWS = 1_000_000
DEFAULT_SAMPLE_ROWS = 5_000
EXPECTED_DATA_ORIGIN = "SYNTHETIC_PIPELINE_TEST"


EXPECTED_SCHEMAS: dict[str, list[str]] = {
    "01_Landslide_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Landslide_Type", "Trigger", "Slope_deg", "Elevation_m",
        "Lithology", "Soil_Moisture_pct", "Rain_24h_mm",
        "Rain_72h_mm", "Road_Distance_m", "River_Distance_m",
        "Susceptibility_Class", "Event_Label",
        "Source_Code", "Data_Origin",
    ],
    "02_Flood_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Flood_Type", "Elevation_m", "River_Distance_m",
        "Rain_24h_mm", "Rain_7d_mm", "Drainage_Density_km_km2",
        "Water_Level_Index", "Flood_Depth_m", "Inundation_pct",
        "Settlement_Distance_m", "Flood_Severity", "Flood_Label",
        "Source_Code", "Data_Origin",
    ],
    "03_Weather_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Temperature_C", "Feels_Like_C", "Humidity_pct",
        "Pressure_hPa", "Wind_Speed_ms", "Wind_Direction_deg",
        "Cloud_Cover_pct", "Visibility_km", "Dew_Point_C",
        "Precip_Probability_pct", "Weather_Condition",
        "Source_Code", "Data_Origin",
    ],
    "04_Rainfall_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Rain_Daily_mm", "Rain_3d_mm", "Rain_7d_mm",
        "Rain_30d_mm", "Rainy_Hours", "Mean_Intensity_mm_h",
        "Antecedent_Rain_mm", "Rainfall_Anomaly_pct",
        "Rain_Class", "Extreme_Rain_Label",
        "Source_Code", "Data_Origin",
    ],
    "05_Terrain_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Elevation_m", "Slope_deg", "Aspect_deg",
        "Profile_Curvature", "Plan_Curvature", "TWI",
        "Terrain_Ruggedness_Index", "Ruggedness",
        "Drainage_Distance_m", "Valley_Depth_m",
        "Terrain_Class", "Source_Code", "Data_Origin",
    ],
    "06_Soil_Geology_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Soil_Group", "Lithology", "Clay_pct", "Sand_pct",
        "Silt_pct", "Bulk_Density_g_cm3", "pH",
        "Organic_Carbon_pct", "Permeability_Index",
        "Fault_Distance_m", "Weathering_Class",
        "Source_Code", "Data_Origin",
    ],
    "07_LULC_Vegetation_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Land_Cover_Class", "NDVI", "EVI", "NDWI",
        "Canopy_Cover_pct", "Builtup_Fraction_pct",
        "Bare_Soil_pct", "Vegetation_Change_pct",
        "Change_Type", "Source_Code", "Data_Origin",
    ],
    "08_Water_Drainage_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Water_Feature_Type", "Stream_Order",
        "River_Distance_m", "Drainage_Density_km_km2",
        "Watershed_ID", "Flow_Accumulation",
        "Water_Occurrence_pct", "Wetness_Index",
        "Bank_Erosion_Risk", "Floodplain_Flag",
        "Source_Code", "Data_Origin",
    ],
    "09_Roads_Infrastructure_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Road_Class", "Road_Distance_m", "Bridge_Distance_m",
        "Settlement_Distance_m", "Building_Density_per_km2",
        "Slope_deg", "Rain_24h_mm", "Cut_Slope_Height_m",
        "Drainage_Condition", "Road_Condition", "Blockage_Risk",
        "Source_Code", "Data_Origin",
    ],
    "10_Historical_Disaster_Risk_10L.xlsx": [
        "Record_ID", "State", "District", "Location_Ref",
        "Latitude", "Longitude", "Observation_Date",
        "Disaster_Type", "Severity_1_5", "Affected_Area_km2",
        "Affected_Population_Est", "Infrastructure_Damage_Index",
        "Recurrence_Class", "Season", "Hazard_Combination",
        "Risk_Score_0_100", "Risk_Class", "High_Risk_Label",
        "Source_Code", "Data_Origin",
    ],
}


RANGE_RULES: dict[str, tuple[float, float]] = {
    "Latitude": (-90.0, 90.0),
    "Longitude": (-180.0, 180.0),
    "Slope_deg": (0.0, 90.0),
    "Humidity_pct": (0.0, 100.0),
    "Cloud_Cover_pct": (0.0, 100.0),
    "Precip_Probability_pct": (0.0, 100.0),
    "Inundation_pct": (0.0, 100.0),
    "Soil_Moisture_pct": (0.0, 100.0),
    "Canopy_Cover_pct": (0.0, 100.0),
    "Builtup_Fraction_pct": (0.0, 100.0),
    "Bare_Soil_pct": (0.0, 100.0),
    "Water_Occurrence_pct": (0.0, 100.0),
    "Risk_Score_0_100": (0.0, 100.0),
    "Severity_1_5": (1.0, 5.0),
    "pH": (0.0, 14.0),
    "NDVI": (-1.0, 1.0),
    "EVI": (-1.0, 1.0),
    "NDWI": (-1.0, 1.0),
    "Wind_Direction_deg": (0.0, 360.0),
    "Rainy_Hours": (0.0, 24.0),
}


NON_NEGATIVE_FIELDS = {
    "Elevation_m",
    "Rain_24h_mm",
    "Rain_72h_mm",
    "Rain_7d_mm",
    "Rain_3d_mm",
    "Rain_30d_mm",
    "Rain_Daily_mm",
    "Road_Distance_m",
    "River_Distance_m",
    "Settlement_Distance_m",
    "Bridge_Distance_m",
    "Drainage_Distance_m",
    "Fault_Distance_m",
    "Flood_Depth_m",
    "Visibility_km",
    "Wind_Speed_ms",
    "Mean_Intensity_mm_h",
    "Antecedent_Rain_mm",
    "Drainage_Density_km_km2",
    "Flow_Accumulation",
    "Affected_Area_km2",
    "Affected_Population_Est",
    "Building_Density_per_km2",
    "Cut_Slope_Height_m",
}


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _to_float(value: Any) -> float | None:
    if _is_blank(value):
        return None

    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def validate_workbook(
    path: str | Path,
    *,
    sample_rows: int = DEFAULT_SAMPLE_ROWS,
) -> dict[str, Any]:
    path = Path(path)

    expected_headers = EXPECTED_SCHEMAS.get(path.name)

    if expected_headers is None:
        return {
            "file_name": path.name,
            "status": "FAIL",
            "errors": ["Unknown dataset filename/schema."],
            "warnings": [],
        }

    errors: list[str] = []
    warnings: list[str] = []

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    if "Data" not in wb.sheetnames:
        wb.close()

        return {
            "file_name": path.name,
            "status": "FAIL",
            "errors": ["Required sheet 'Data' is missing."],
            "warnings": [],
        }

    ws = wb["Data"]

    actual_rows = max(ws.max_row - 1, 0)
    actual_columns = ws.max_column

    headers = [
        cell.value
        for cell in next(
            ws.iter_rows(
                min_row=1,
                max_row=1,
            )
        )
    ]

    if actual_rows != EXPECTED_ROWS:
        errors.append(
            f"Expected {EXPECTED_ROWS} data rows, found {actual_rows}."
        )

    if headers != expected_headers:
        errors.append("Workbook schema/header order does not match expected schema.")

    if actual_columns != len(expected_headers):
        errors.append(
            f"Expected {len(expected_headers)} columns, "
            f"found {actual_columns}."
        )

    header_index = {
        str(name): index
        for index, name in enumerate(headers)
        if name is not None
    }

    sampled = 0
    duplicate_ids = 0
    blank_ids = 0
    blank_source_codes = 0
    blank_states = 0
    blank_districts = 0
    invalid_coordinates = 0
    invalid_data_origin = 0
    blank_dates = 0

    seen_ids: set[str] = set()
    origin_counts: Counter[str] = Counter()

    range_violations: Counter[str] = Counter()
    non_negative_violations: Counter[str] = Counter()
    non_numeric_values: Counter[str] = Counter()

    max_row_to_read = min(
        ws.max_row,
        sample_rows + 1,
    )

    for row in ws.iter_rows(
        min_row=2,
        max_row=max_row_to_read,
        values_only=True,
    ):
        sampled += 1

        def get(column: str) -> Any:
            index = header_index.get(column)

            if index is None or index >= len(row):
                return None

            return row[index]

        record_id = get("Record_ID")

        if _is_blank(record_id):
            blank_ids += 1
        else:
            record_key = str(record_id).strip()

            if record_key in seen_ids:
                duplicate_ids += 1
            else:
                seen_ids.add(record_key)

        if _is_blank(get("State")):
            blank_states += 1

        if _is_blank(get("District")):
            blank_districts += 1

        if _is_blank(get("Observation_Date")):
            blank_dates += 1

        if _is_blank(get("Source_Code")):
            blank_source_codes += 1

        origin = str(get("Data_Origin") or "").strip()

        origin_counts[origin] += 1

        if origin != EXPECTED_DATA_ORIGIN:
            invalid_data_origin += 1

        latitude = _to_float(get("Latitude"))
        longitude = _to_float(get("Longitude"))

        if (
            latitude is None
            or longitude is None
            or not (-90 <= latitude <= 90)
            or not (-180 <= longitude <= 180)
        ):
            invalid_coordinates += 1

        for column, (minimum, maximum) in RANGE_RULES.items():
            if column not in header_index:
                continue

            value = get(column)

            if _is_blank(value):
                continue

            number = _to_float(value)

            if number is None:
                non_numeric_values[column] += 1
                continue

            if not minimum <= number <= maximum:
                range_violations[column] += 1

        for column in NON_NEGATIVE_FIELDS:
            if column not in header_index:
                continue

            value = get(column)

            if _is_blank(value):
                continue

            number = _to_float(value)

            if number is None:
                non_numeric_values[column] += 1
                continue

            if number < 0:
                non_negative_violations[column] += 1

    wb.close()

    if sampled < min(sample_rows, EXPECTED_ROWS):
        warnings.append(
            f"Only {sampled} sample rows were available for content validation."
        )

    if duplicate_ids:
        errors.append(
            f"{duplicate_ids} duplicate Record_ID values found in sampled rows."
        )

    if blank_ids:
        errors.append(
            f"{blank_ids} blank Record_ID values found in sampled rows."
        )

    if invalid_coordinates:
        errors.append(
            f"{invalid_coordinates} invalid coordinates found in sampled rows."
        )

    if invalid_data_origin:
        errors.append(
            f"{invalid_data_origin} sampled rows do not use "
            f"Data_Origin={EXPECTED_DATA_ORIGIN}."
        )

    if blank_source_codes:
        warnings.append(
            f"{blank_source_codes} blank Source_Code values found in sampled rows."
        )

    if blank_states:
        warnings.append(
            f"{blank_states} blank State values found in sampled rows."
        )

    if blank_districts:
        warnings.append(
            f"{blank_districts} blank District values found in sampled rows."
        )

    if blank_dates:
        warnings.append(
            f"{blank_dates} blank Observation_Date values found in sampled rows."
        )

    if range_violations:
        errors.append(
            "Range violations detected: "
            + ", ".join(
                f"{column}={count}"
                for column, count in sorted(range_violations.items())
            )
        )

    if non_negative_violations:
        errors.append(
            "Negative values detected in non-negative fields: "
            + ", ".join(
                f"{column}={count}"
                for column, count in sorted(
                    non_negative_violations.items()
                )
            )
        )

    if non_numeric_values:
        warnings.append(
            "Non-numeric values detected in numeric fields: "
            + ", ".join(
                f"{column}={count}"
                for column, count in sorted(non_numeric_values.items())
            )
        )

    status = (
        "FAIL"
        if errors
        else "PASS_WITH_WARNINGS"
        if warnings
        else "PASS"
    )

    return {
        "file_name": path.name,
        "classification": "synthetic_pipeline_test",
        "expected_rows": EXPECTED_ROWS,
        "actual_rows": actual_rows,
        "columns": actual_columns,
        "sample_rows_checked": sampled,
        "data_origin_counts": dict(origin_counts),
        "duplicate_ids_in_sample": duplicate_ids,
        "invalid_coordinates_in_sample": invalid_coordinates,
        "range_violations": dict(range_violations),
        "non_negative_violations": dict(non_negative_violations),
        "errors": errors,
        "warnings": warnings,
        "real_world_evidence_eligible": False,
        "status": status,
    }


def validate_synthetic_10m_folder(
    root: str | Path,
    *,
    sample_rows: int = DEFAULT_SAMPLE_ROWS,
) -> dict[str, Any]:
    root = Path(root)

    reports: list[dict[str, Any]] = []

    for file_name in EXPECTED_SCHEMAS:
        path = root / file_name

        if not path.exists():
            reports.append(
                {
                    "file_name": file_name,
                    "status": "FAIL",
                    "errors": ["Expected workbook is missing."],
                    "warnings": [],
                }
            )
            continue

        reports.append(
            validate_workbook(
                path,
                sample_rows=sample_rows,
            )
        )

    failed = [
        report
        for report in reports
        if report["status"] == "FAIL"
    ]

    warning_reports = [
        report
        for report in reports
        if report["status"] == "PASS_WITH_WARNINGS"
    ]

    total_rows = sum(
        int(report.get("actual_rows", 0))
        for report in reports
    )

    status = (
        "FAIL"
        if failed
        else "PASS_WITH_WARNINGS"
        if warning_reports
        else "PASS"
    )

    return {
        "classification": "synthetic_10m_collection",
        "dataset_count": len(reports),
        "total_rows": total_rows,
        "reports": reports,
        "failed_dataset_count": len(failed),
        "warning_dataset_count": len(warning_reports),
        "real_world_evidence_eligible": False,
        "status": status,
    }
